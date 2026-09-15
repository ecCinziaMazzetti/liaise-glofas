#!/usr/bin/env python3
"""Build docs/eclandpy_liaise_report.docx and docs/eclandpy_liaise_slides.pptx from the figures
in docs/figures/ and the numbers in the two diagnostics JSONs + the GRDC results. Re-run after
any rerun of the chain; python-docx / python-pptx (system python3)."""
from __future__ import annotations

import json
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.util import Emu, Inches, Pt as PPt

D = Path("/perm/pad/liaise-ecland/docs")
FIG = D / "figures"
ACCENT = (0x2A, 0x6F, 0xB0)
INK = (0x1C, 0x1B, 0x19)
MUTED = (0x5A, 0x57, 0x50)

# ------------------------------------------------------------------ numbers (from the run)
e = json.load(open("/perm/pad/liaise_discharge_compare/eclandpy_run_diagnostics.json"))["annual"]
c = json.load(open("/perm/pad/liaise_discharge_compare/control_run_diagnostics.json"))["annual"]
ys = sorted(set(e) & set(c), key=int)
def arr(d, k, f=1.0): return np.array([d[y][k] for y in ys]) * f
def stat(k, fe=1.0, fc=1.0, dec=1):
    a, b = arr(e, k, fe), arr(c, k, fc)
    return f"{a.mean():.{dec}f}", f"{b.mean():.{dec}f}", f"{a.mean()-b.mean():+.{dec}f}", f"{np.corrcoef(a, b)[0, 1]:.3f}"
LAND = [
    ("Precipitation, mm/yr", *stat("precip_mm")),
    ("Evapotranspiration, mm/yr", *stat("evap_mm")),
    ("Total runoff −(Qs+Qsb), mm/yr", *stat("runoff_mm", 1, -1)),
    ("Surface T (AvgSurfT vs T2m), °C", *stat("t2m_mean_c", dec=2)),
    ("Root-zone soil moisture, kg/m²", *stat("rootmoist_mean")),
]
r = json.load(open("/perm/pad/liaise_discharge_compare/skill_benchmark_results_eclandpy.json"))
r = [x for x in r if x["station"] != "RIO GUADALOPE, CASPE"]
def med(mo, k): return np.median([x[k] for x in r if x["model"] == mo and x[k] is not None and np.isfinite(x[k])])
keys = {(x["year"], x["station"]) for x in r}
wins = sum(1 for kk in keys for g in [next((x for x in r if x["model"] == "GPU" and (x["year"], x["station"]) == kk), None)]
           for f in [next((x for x in r if x["model"] == "Fortran" and (x["year"], x["station"]) == kk), None)]
           if g and f and np.isfinite(g["kge"]) and np.isfinite(f["kge"]) and g["kge"] > f["kge"])
GRDC = [("median KGE", f"{med('GPU','kge'):.3f}", f"{med('Fortran','kge'):.3f}"),
        ("median NSE", f"{med('GPU','nse'):.3f}", f"{med('Fortran','nse'):.3f}"),
        ("median correlation r", f"{med('GPU','r'):.3f}", f"{med('Fortran','r'):.3f}"),
        ("median PBIAS, %", f"{med('GPU','pbias_pct'):.1f}", f"{med('Fortran','pbias_pct'):.1f}"),
        ("station-years won on KGE", f"{wins} of {len(keys)}", f"{len(keys)-wins} of {len(keys)}")]
PERF = [("Fortran ecLand CY50R1, 1 node, 4 OpenMP threads", "2.56 min", "~8.8 ms", "1 h 35 m (measured)"),
        ("eclandpy, gt:cpu_kfirst, 1 CPU core", "30.5 min", "~109 ms", "~19 h"),
        ("eclandpy, gt:gpu, 1× NVIDIA A100", "10.6 min", "31 ms", "6.5 h (measured, incl. resume)")]

# ------------------------------------------------------------------ shared prose
SUMMARY = (
    "We ran the full LIAISE land–river chain entirely in Python: eclandpy — a GT4Py port of ecLand "
    "with cy50r1 physics, executed on one NVIDIA A100 — for the land surface over 37 years (1988–2024) "
    "on the 16×23 LIAISE grid, and CaMa-Flood-GPU (Kang, Yin & Yamazaki 2026) for river routing. "
    "Against the Fortran ecLand CY50R1 control run with identical forcing and ancillaries, surface "
    "temperature is indistinguishable (+0.05 °C, r = 0.99 over 37 years), evapotranspiration is 2 % low, "
    "and total runoff 21 % low — the one substantive physics gap, traced to a regionally coherent "
    "deep-soil-moisture divergence that domain means hide. Routed to the six GRDC gauges of the domain, "
    "the Python chain scores marginally better than the Fortran chain (median KGE −0.135 vs −0.155, "
    "r 0.45 vs 0.34, ahead in 15 of 25 station-years). The year-by-year restart mechanism built for eclandpy "
    "reproduces the Fortran restart state after ten chained years to within 0.1 K in soil temperature and "
    "1–2 % in soil moisture. On this small domain the GPU is 2.9× faster than one CPU core but still ~4× "
    "slower than Fortran on four cores: 368 columns cannot fill an A100, so the port pays off at scale."
)
STRATEGY = [
    ("Reuse, don't re-derive", "eclandpy drives the Fortran-twin-validated GT4Py kernels of ecland_porting "
     "(C. Kühnlein) unmodified through their own offline driver; eclandpy's own code is the I/O layer, the "
     "adapter, and what the port lacks."),
    ("Close the cy48r1→cy50r1 gap surgically", "Diffing two clean Fortran checkouts across every module on the "
     "real execution path found that most cy50r1 physics was already in the 'cy48r1' port; three genuine gaps "
     "were ported (RTF2 freezing curve, snow-cover fraction, vegetation tables) and the seven experimental "
     "hydrology flags kept off on both sides — a plain cy50r1 vs plain cy50r1 comparison."),
    ("Validate on points before grids", "PLUMBER2: 42 then all 170 flux-tower sites, scored against FLUXNET; "
     "Qle/Qh match Fortran within noise, NEE is the known weaker variable. Only then the first multi-point run "
     "ever attempted with this core — the LIAISE grid — and its 5-day check against Fortran."),
    ("Build what the offline driver lacks", "A restart chain (ecland_porting has none), diagnostic recorders, "
     "lat/lon on outputs so Python files are drop-ins for the Fortran-schema tools, and a coupling layer "
     "(eclandpy.cmfgpu) that mirrors cnt41s.F90's hand-over: −(Qs+Qsb), the weights, the advance loop."),
    ("Keep the reference run as the arbiter", "Every step is checked against the Fortran control: initial "
     "state, budgets, restart state, annual diagnostics, gauge skill. Where the Python chain is better or worse, "
     "we say so with the number."),
]
LESSONS = [
    "GPU needs its own environment: cupy-cuda12x 13.6 (14.x forces numpy≥2), nanobind 1.9.2, gcc 11 as nvcc host "
    "(nvcc rejects >11), CUDA 12.6, LD_LIBRARY_PATH for GLIBCXX_3.4.26, GT4PY_EXTRA_COMPILE_ARGS unset, ≥64 GB for "
    "the nvcc compile of the largest stencil, SLURM qos ng.",
    "gt4py compiles stencils lazily and re-reads their source from disk: editing model source under a live run "
    "recompiles and crashes it (cost us a 20-hour and a 2-hour run). Never touch the source tree while a job runs.",
    "stencil_validation went public during this work: the pmap-ifs shim eclandpy used lacked freeze() and "
    "get_stencil_id; the real package rejects the registry re-registration Phase 4a relied on and re-registers "
    "per instantiation. All three now handled; both backends pass single- and multi-driver tests.",
    "Sign conventions cost real effort: Qsb is a negative soil-column loss, so total runoff is −(Qs+Qsb); "
    "Qs−Qsb double-counts surface runoff and passed every plausibility check until compared with Fortran discharge.",
]
GAPS = [
    "Deep-layer (1.9 m) soil moisture diverges regionally after ~2 years — eclandpy drier on the Duero plateau, "
    "wetter on the Mediterranean and Cantabrian coasts; 35 of 235 cells by >100 kg/m². The runoff and "
    "root-zone offsets follow from it. Next: locate when and where it develops in the L4 time series.",
    "No T2m/D2m diagnostic in eclandpy (surfpp_ctl_mod's per-tile 2 m scheme is unported); AvgSurfT stands in, "
    "which also explains the larger seasonal amplitude in the climatology.",
    "Two Fortran coupling switches are unported and off in the reference: LWEVAP (lake-tile potential "
    "evaporation extracted from floodplain storage — CaMa-Flood-GPU has no evaporation sink) and LROSPLIT "
    "(surface/sub-surface passed separately — equivalent while LGDWDLY is off).",
    "The land-ice tile is dead code in the port (InputBinder hardcodes pcil = 0), harmless on LIAISE; "
    "NEE remains the weakest PLUMBER2 variable.",
    "Performance case must be made at scale: at 368 columns per-step cost is launch latency; CaMa-Flood-GPU's "
    "own benchmarks run at 17,675×N columns. Batching the 170 PLUMBER2 sites into one domain is the natural test.",
]
REPOS = [
    ("github.com/gpbalsamo/eclandpy", "the land model + eclandpy.cmfgpu coupling layer (this work: 4e35128 and later)"),
    ("github.com/gpbalsamo/liaise-ecland", "LIAISE configuration, Fortran control, eclandpy_bridge/ (93fae86 and later), the archived CaMa-Flood-GPU drivers"),
    ("github.com/gpbalsamo/ecland-porting, branch cy50r1", "the three cy50r1 physics ports on top of C. Kühnlein's main (merged 29e6152)"),
    ("github.com/Kshy0/CaMa-Flood-GPU + Kshy0/hydroforge", "the routing model, unmodified"),
    ("github.com/stubbiali/stencil-validation (public since 2026-09-14)", "the stencil framework at ecland-porting's pinned commit"),
]

# ------------------------------------------------------------------ Word
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
for s in doc.sections: s.left_margin = s.right_margin = Cm(2.2)
def H(t, lvl=1):
    h = doc.add_heading(t, lvl)
    for run in h.runs: run.font.color.rgb = RGBColor(*INK)
    return h
def P(t, italic=False, size=None):
    p = doc.add_paragraph(); run = p.add_run(t); run.italic = italic
    if size: run.font.size = Pt(size)
    return p
def T(rows, header, widths=None):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Light Grid Accent 1"
    for i, h in enumerate(header): t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): cells[i].text = str(v)
    for row in t.rows:
        for i, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                for run in p.runs: run.font.size = Pt(9.5)
                if i > 0: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph()
def FIGURE(name, caption, width=16):
    doc.add_picture(str(FIG / name), width=Cm(width)); cp = P(caption, italic=True, size=9); cp.alignment = WD_ALIGN_PARAGRAPH.LEFT

t = doc.add_heading("An all-Python ecLand–CaMa-Flood chain: eclandpy on the LIAISE domain", 0)
P("Porting strategy, validation against the Fortran ecLand CY50R1 control, and results, 1988–2024. "
  "G. Balsamo, ECMWF, September 2026 — for N. Wedi and C. Kühnlein.", italic=True)
H("Summary"); P(SUMMARY)
H("1. Porting strategy")
P("The aim was a standalone Python land-surface model behaving like ecLand — same namelists, same NetCDF I/O, "
  "same physics — validated at every step against the Fortran, and coupled to the GPU reimplementation of "
  "CaMa-Flood so that the whole land–river chain runs in Python. Five principles carried the work:")
for ttl, txt in STRATEGY:
    p = doc.add_paragraph(style="List Bullet"); rr = p.add_run(ttl + ". "); rr.bold = True; p.add_run(txt)
H("2. What was built", 1)
P("eclandpy (repo gpbalsamo/eclandpy): the I/O layer for plumber2-ecland and LIAISE file conventions; the adapter "
  "onto ecland_porting's OfflineDriver; Phase 4a (six of the seven cy50r1 hydrology flags as compile-time-gated "
  "overrides, default off); a year-by-year restart chain; and eclandpy.cmfgpu, the coupling layer to CaMa-Flood-GPU "
  "— runoff hand-over, inpmat→npz weights, regional subsetting, the chained multi-year driver and discharge export — "
  "each a CLI, domain-agnostic, with LIAISE as the worked example.")
P("ecland-porting (fork, branch cy50r1): three cy50r1 physics ports on top of C. Kühnlein's cleaned-up main "
  "(merged 2026-09-14): the RTF2 freezing-curve constant, the LESN09 snow-cover-fraction reformulation (a standing "
  "bug — the port matched neither cycle's default), and the vegetation tables.")
P("liaise-ecland: the LIAISE configuration of all of it, the Fortran control run, and the archived CaMa-Flood-GPU "
  "driver scripts that produced every documented number.")
H("3. Environments and lessons", 1)
for l in LESSONS: doc.add_paragraph(l, style="List Bullet")
H("4. Results: land surface vs the Fortran control", 1)
P(f"Domain means over the 235 active land points, all {len(ys)} common years. Precipitation identical confirms "
  "identical forcing; the runoff and root-zone deficits are the same signal.")
T(LAND, ["", "eclandpy (GPU)", "Fortran", "bias", "r"])
FIGURE("annual_series.png", "Figure 1. Annual domain means, eclandpy (blue) vs the Fortran control (grey). Both start from the same soilinit in 1988; the runoff and root-zone offsets establish within two years and then persist as an offset, not a drift.")
FIGURE("seasonal_cycle.png", "Figure 2. Climatology over 37 years. The larger temperature amplitude in eclandpy is the skin-vs-2 m distinction (AvgSurfT vs T2m), not a physics difference.")
H("5. The restart chain, checked against the Fortran restart", 1)
P("ecland_porting has no restart; eclandpy pickles the driver's carried state (OfflineState, generically over its "
  "dataclass fields, plus the carbon accumulators) at each year boundary and re-injects it into a fresh driver. "
  "Checked at 1998-01-01 after ten chained years against the Fortran run's own restart_in.nc: identical at 1988 "
  "t=0 (the surfinit conversion is exact); AvgSurfT 276.96 vs 277.02 K; soil temperature within 0.1 K and soil "
  "moisture within 1–2 % in all four layers; one clean timestep across the 1997→1998 boundary. Two real "
  "differences: SWE 0.83 vs 0.56 and canopy water 0.22 vs 0.14 kg/m² (small, snow and interception), and the "
  "deep-layer divergence of §7.")
H("6. Results: river discharge vs GRDC gauges", 1)
P("Runoff −(Qs+Qsb) routed by CaMa-Flood-GPU with river storage carried across all 37 years, scored on daily "
  "discharge at the six LIAISE gauges over the five benchmark years (Guadalope at Caspe excluded as a regulated "
  "river), against the Fortran ecLand + CaMa-Flood control chain.")
T(GRDC, ["", "eclandpy → CaMa-Flood-GPU", "Fortran chain"])
FIGURE("grdc_kge.png", "Figure 3. Median KGE per gauge. Both chains are poor in absolute terms on this domain; the Python chain is marginally better correlated and drier.")
H("7. Known gaps and next steps", 1)
for g in GAPS: doc.add_paragraph(g, style="List Bullet")
H("8. Performance", 1)
T(PERF, ["configuration", "per year", "per step", "37 years"])
FIGURE("performance.png", "Figure 4. Wall-clock per simulated year on the 368-column LIAISE grid.", 12)
H("9. Reproducing this work", 1)
P("liaise-ecland/eclandpy_bridge/README.md gives the four commands (prepare inputs, land run, routing chain, "
  "dashboard) and both environments; eclandpy/README.md the coupling CLIs. Dashboards: "
  "sites.ecmwf.int/pad/liaise/eclandpy/ beside /pad/liaise/control/; sites.ecmwf.int/pad/plumber2/dashboards/ for PLUMBER2.")
T(REPOS, ["repository", "role"])
doc.save(D / "eclandpy_liaise_report.docx")

# ------------------------------------------------------------------ PowerPoint
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
def slide(title, sub=None):
    s = prs.slides.add_slide(BLANK)
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9)); tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title; p.font.size = PPt(30); p.font.bold = True; p.font.color.rgb = PRGB(*INK)
    if sub:
        p2 = tf.add_paragraph(); p2.text = sub; p2.font.size = PPt(15); p2.font.color.rgb = PRGB(*MUTED)
    ln = s.shapes.add_shape(1, Inches(0.6), Inches(1.32), Inches(1.2), Emu(38000)); ln.fill.solid(); ln.fill.fore_color.rgb = PRGB(*ACCENT); ln.line.fill.background()
    return s
def bullets(s, items, x=0.6, y=1.6, w=12.1, h=5.4, size=16, bold_lead=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if bold_lead and isinstance(it, tuple):
            r1 = p.add_run(); r1.text = it[0] + "  "; r1.font.bold = True; r1.font.size = PPt(size); r1.font.color.rgb = PRGB(*ACCENT)
            r2 = p.add_run(); r2.text = it[1]; r2.font.size = PPt(size); r2.font.color.rgb = PRGB(*INK)
        else:
            p.text = "•  " + it; p.font.size = PPt(size); p.font.color.rgb = PRGB(*INK)
        p.space_after = PPt(9)
def table(s, rows, header, x=0.6, y=1.6, w=12.1, size=14, rh=0.42):
    t = s.shapes.add_table(len(rows) + 1, len(header), Inches(x), Inches(y), Inches(w), Inches(rh * (len(rows) + 1))).table
    for j, h in enumerate(header): t.cell(0, j).text = h
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row): t.cell(i, j).text = str(v)
    for i in range(len(rows) + 1):
        for j in range(len(header)):
            for p in t.cell(i, j).text_frame.paragraphs:
                p.font.size = PPt(size); p.font.bold = i == 0
def picture(s, name, x=0.6, y=1.55, w=12.1):
    s.shapes.add_picture(str(FIG / name), Inches(x), Inches(y), width=Inches(w))

s = slide("An all-Python ecLand–CaMa-Flood chain", "eclandpy on the LIAISE domain, 1988–2024, validated against the Fortran ecLand CY50R1 control")
bullets(s, ["G. Balsamo — ECMWF — September 2026", "with C. Kühnlein (ecland-porting) and N. Wedi (pmap-ifs)",
            "Land: eclandpy (GT4Py port of ecLand, cy50r1 physics) on one NVIDIA A100",
            "River: CaMa-Flood-GPU (Kang, Yin & Yamazaki 2026, PyTorch/Triton)"], y=2.2, size=18)
s = slide("In one slide", "what was done, and what it showed")
bullets(s, [
    "37 years on the 16×23 LIAISE grid: the first multi-point run of this Python core, then routed through CaMa-Flood-GPU",
    "Surface temperature indistinguishable from Fortran (+0.05 °C, r = 0.99); evaporation 2 % low; runoff 21 % low",
    "Restart chain reproduces the Fortran restart state after 10 years: soil T within 0.1 K, soil moisture within 1–2 %",
    "At the 6 GRDC gauges the Python chain edges the Fortran chain: median KGE −0.135 vs −0.155, r 0.45 vs 0.34, 15/25 station-years",
    "One real physics gap: a regionally coherent deep-soil-moisture divergence that domain means hide",
    "GPU 2.9× faster than one CPU core, ~4× slower than Fortran on 4 cores — the domain is too small to fill an A100",
], size=17)
s = slide("Porting strategy", "five principles")
bullets(s, STRATEGY, size=15, bold_lead=True)
s = slide("Architecture", "same structure as the Fortran, different packaging")
bullets(s, [
    ("ecLand Fortran", "vendors CaMa-Flood (src/surf/cmflood) and carries the coupling in cnt41s.F90 — one executable"),
    ("eclandpy", "drives ecland_porting's GT4Py kernels unmodified; adds I/O, adapter, Phase 4a flags, a restart chain"),
    ("eclandpy.cmfgpu", "the coupling layer only — −(Qs+Qsb) hand-over, inpmat→npz weights, regional subset, chained driver, discharge — CLIs, domain-agnostic"),
    ("CaMa-Flood-GPU", "stays an external package: its torch/triton stack cannot share a process with gt4py 1.1.7 / numpy<2"),
    ("liaise-ecland", "is configuration — which parameters.nc, mapping, gauges — the way namelist/input_cmf configures LECMF1WAY"),
], size=15, bold_lead=True)
s = slide("Environments and lessons learned"); bullets(s, LESSONS, size=15)
s = slide("Land surface vs the Fortran control", f"domain means over 235 land points, {len(ys)} years, identical forcing")
table(s, LAND, ["", "eclandpy (GPU)", "Fortran", "bias", "r"], w=12.1, size=15)
bullets(s, ["Runoff and root-zone deficits are one signal: water stays in the deep soil rather than leaving the column",
            "Fortran's o_wat does not close its own budget (−60 to −70 mm/yr), so part of the runoff gap may be diagnostic"], y=4.6, size=15)
s = slide("Annual series, 1988–2024"); picture(s, "annual_series.png", x=2.07, w=9.2)
s = slide("Seasonal cycle", "the larger temperature amplitude is skin (AvgSurfT) vs 2 m — eclandpy has no T2m diagnostic"); picture(s, "seasonal_cycle.png", x=0.9, w=11.5)
s = slide("The restart chain, checked against the Fortran restart", "state at 1998-01-01 after ten chained years")
bullets(s, ["Identical to Fortran at 1988 t=0 — the surfinit conversion is exact",
            "AvgSurfT 276.96 vs 277.02 K; soil temperature within 0.1 K in all four layers",
            "Soil moisture within 1–2 % (domain mean) in all four layers; one clean timestep across the 1997→1998 boundary",
            "Small real differences: SWE 0.83 vs 0.56, canopy water 0.22 vs 0.14 kg/m²",
            "Deep layer (1.9 m): 35 of 235 cells differ by >100 kg/m² — drier on the Duero plateau, wetter on the coasts; signs balance, so means hide it"], size=16)
s = slide("River discharge vs GRDC observations", "6 LIAISE gauges, 5 benchmark years, 25 station-years")
table(s, GRDC, ["", "eclandpy → CaMa-Flood-GPU", "Fortran chain"], w=8, size=13, rh=0.34)
picture(s, "grdc_kge.png", x=1.87, y=3.7, w=9.6)
s = slide("Performance", "per simulated year, 368-column LIAISE grid")
table(s, PERF, ["configuration", "per year", "per step", "37 years"], size=14)
picture(s, "performance.png", x=3.67, y=3.6, w=6.0)
bullets(s, ["Per-step cost at 368 columns is kernel-launch latency; CaMa-Flood-GPU's own benchmarks run at 17,675×N columns — the case for the port is at scale"], y=6.45, h=0.9, size=14)
s = slide("Known gaps and next steps"); bullets(s, GAPS, size=14)
s = slide("Reproduce it", "four commands, two environments")
bullets(s, ["liaise-ecland/eclandpy_bridge/README.md — prepare inputs → land run (CPU or A100) → routing chain → dashboard",
            "eclandpy/README.md 'CaMa-Flood-GPU coupling' — the eclandpy.cmfgpu CLIs and the GPU environment",
            "Dashboards: sites.ecmwf.int/pad/liaise/eclandpy/ (beside /control/), sites.ecmwf.int/pad/plumber2/dashboards/"] +
           [f"{a} — {b}" for a, b in REPOS], size=14)
prs.save(D / "eclandpy_liaise_slides.pptx")
print("wrote", D / "eclandpy_liaise_report.docx", "and", D / "eclandpy_liaise_slides.pptx")
